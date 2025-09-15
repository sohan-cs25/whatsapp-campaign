from django.shortcuts import render
from django.http import HttpResponse, FileResponse, Http404, StreamingHttpResponse
from django.core.exceptions import ValidationError
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import csrf_exempt
from rest_framework import status, viewsets, filters
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.pagination import PageNumberPagination
from django.db.models import Count, Q, Sum
from django.utils import timezone
import os
import logging
import mimetypes
import json
import time

from .whatsapp_service import OrderWhatsAppService

from .models import ChatFile, ParsedChatFile, ValidatedFile, Order, WebhookEvent
from .serializers import (
    ChatFileSerializer, ChatFileListSerializer,
    ParsedChatFileSerializer, ParsedChatFileListSerializer,
    ValidatedFileSerializer, ValidatedFileListSerializer,
    OrderSerializer, OrderListSerializer,
    WebhookEventSerializer, FileUploadProgressSerializer,
    ProcessingStatsSerializer
)

logger = logging.getLogger(__name__)


class StandardResultsSetPagination(PageNumberPagination):
    page_size = 20
    page_size_query_param = 'page_size'
    max_page_size = 100


class ChatFileViewSet(viewsets.ModelViewSet):
    """
    ViewSet for ChatFile CRUD operations
    Handles WhatsApp .txt file uploads
    """
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['filename']
    ordering_fields = ['uploaded_at', 'is_processed']
    ordering = ['-uploaded_at']
    pagination_class = StandardResultsSetPagination

    def get_queryset(self):
        return ChatFile.objects.filter(user=self.request.user)

    def get_serializer_class(self):
        if self.action == 'list':
            return ChatFileListSerializer
        return ChatFileSerializer

    def create(self, request, *args, **kwargs):
        """Upload a new chat file"""
        try:
            if 'filepath' not in request.FILES:
                return Response({
                    'error': 'No file provided. Please upload a file.'
                }, status=status.HTTP_400_BAD_REQUEST)

            uploaded_file = request.FILES['filepath']

            # Validate file type
            if not uploaded_file.name.endswith(('.txt', '.csv')):
                return Response({
                    'error': 'Invalid file type. Only .txt and .csv files are allowed.'
                }, status=status.HTTP_400_BAD_REQUEST)

            # Validate file size (max 50MB)
            if uploaded_file.size > 50 * 1024 * 1024:
                return Response({
                    'error': 'File too large. Maximum size is 50MB.'
                }, status=status.HTTP_400_BAD_REQUEST)

            # Set filename and filetype
            data = request.data.copy()
            data['filename'] = uploaded_file.name
            data['filetype'] = 'text' if uploaded_file.name.endswith('.txt') else 'csv'

            serializer = self.get_serializer(data=data)
            if serializer.is_valid():
                chat_file = serializer.save()
                logger.info(f"Chat file uploaded: {chat_file.filename} by user {request.user.username}")

                return Response({
                    'success': True,
                    'message': 'File uploaded successfully',
                    'data': serializer.data
                }, status=status.HTTP_201_CREATED)

            return Response({
                'success': False,
                'errors': serializer.errors
            }, status=status.HTTP_400_BAD_REQUEST)

        except Exception as e:
            logger.error(f"Error uploading chat file: {str(e)}")
            return Response({
                'error': 'An error occurred while uploading the file.'
            }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=True, methods=['post'])
    def process(self, request, pk=None):
        """Trigger AI processing of chat file"""
        chat_file = self.get_object()

        if chat_file.is_processed:
            return Response({
                'error': 'File has already been processed'
            }, status=status.HTTP_400_BAD_REQUEST)

        # Trigger Celery task for processing
        from .tasks import process_chat_file
        process_chat_file.delay(str(chat_file.id))

        return Response({
            'success': True,
            'message': 'File processing started',
            'file_id': chat_file.id
        })


class ParsedChatFileViewSet(viewsets.ModelViewSet):
    """
    ViewSet for ParsedChatFile operations
    Handles AI-processed chat files
    """
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.OrderingFilter]
    ordering_fields = ['processed_at', 'total_orders']
    ordering = ['-processed_at']
    pagination_class = StandardResultsSetPagination

    def get_queryset(self):
        return ParsedChatFile.objects.filter(user=self.request.user)

    def get_serializer_class(self):
        if self.action == 'list':
            return ParsedChatFileListSerializer
        return ParsedChatFileSerializer

    @action(detail=True, methods=['get'])
    def download(self, request, pk=None):
        """Download processed chat file"""
        try:
            parsed_file = self.get_object()

            # Build correct file path using ORDERS_MEDIA_ROOT
            from django.conf import settings

            # Handle FileField path - use the name/path, not the FileField object
            if hasattr(parsed_file.processed_file_path, 'name'):
                relative_path = parsed_file.processed_file_path.name
            else:
                relative_path = str(parsed_file.processed_file_path)

            file_path = os.path.join(settings.ORDERS_MEDIA_ROOT, relative_path)

            if not os.path.exists(file_path):
                logger.error(f"File not found at: {file_path}")
                raise Http404("File not found")

            # Update download tracking
            parsed_file.downloaded_at = timezone.now()
            parsed_file.downloaded_by = request.user
            parsed_file.save()

            # Determine content type
            content_type, _ = mimetypes.guess_type(file_path)
            if not content_type:
                content_type = 'application/octet-stream'

            response = FileResponse(
                open(file_path, 'rb'),
                content_type=content_type,
                as_attachment=True,
                filename=parsed_file.file_name
            )

            logger.info(f"Downloaded processed file: {parsed_file.file_name} by user {request.user.username}")
            return response

        except Exception as e:
            logger.error(f"Error downloading processed file: {str(e)}")
            raise Http404("File not found")


class ValidatedFileViewSet(viewsets.ModelViewSet):
    """
    ViewSet for ValidatedFile operations
    Handles human-validated order files
    """
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['file_name']
    ordering_fields = ['uploaded_at', 'orders_extracted']
    ordering = ['-uploaded_at']
    pagination_class = StandardResultsSetPagination

    def get_queryset(self):
        return ValidatedFile.objects.filter(user=self.request.user)

    def get_serializer_class(self):
        if self.action == 'list':
            return ValidatedFileListSerializer
        return ValidatedFileSerializer

    def create(self, request, *args, **kwargs):
        """Upload a validated file"""
        try:
            if 'filepath' not in request.FILES:
                return Response({
                    'error': 'No file provided. Please upload a file.'
                }, status=status.HTTP_400_BAD_REQUEST)

            uploaded_file = request.FILES['filepath']

            # Validate file type
            if not uploaded_file.name.endswith(('.xlsx', '.csv')):
                return Response({
                    'error': 'Invalid file type. Only .xlsx and .csv files are allowed.'
                }, status=status.HTTP_400_BAD_REQUEST)

            # Set file details
            data = request.data.copy()
            data['file_name'] = uploaded_file.name
            data['filetype'] = 'xlsx' if uploaded_file.name.endswith('.xlsx') else 'csv'

            serializer = self.get_serializer(data=data)
            if serializer.is_valid():
                validated_file = serializer.save()
                logger.info(f"Validated file uploaded: {validated_file.file_name} by user {request.user.username}")

                return Response({
                    'success': True,
                    'message': 'Validated file uploaded successfully! You can now extract orders and send messages.',
                    'next_steps': [
                        'Click "Extract Orders" to process the file',
                        'Review the extracted orders',
                        'Click "Send Messages" to start sending WhatsApp messages'
                    ],
                    'data': serializer.data
                }, status=status.HTTP_201_CREATED)

            return Response({
                'success': False,
                'errors': serializer.errors
            }, status=status.HTTP_400_BAD_REQUEST)

        except Exception as e:
            logger.error(f"Error uploading validated file: {str(e)}")
            return Response({
                'error': 'An error occurred while uploading the file.'
            }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=True, methods=['get'])
    def download(self, request, pk=None):
        """Download validated file"""
        try:
            validated_file = self.get_object()
            file_path = validated_file.filepath.path

            if not os.path.exists(file_path):
                raise Http404("File not found")

            # Determine content type
            content_type, _ = mimetypes.guess_type(file_path)
            if not content_type:
                content_type = 'application/octet-stream'

            response = FileResponse(
                open(file_path, 'rb'),
                content_type=content_type,
                as_attachment=True,
                filename=validated_file.file_name
            )

            logger.info(f"Downloaded validated file: {validated_file.file_name} by user {request.user.username}")
            return response

        except Exception as e:
            logger.error(f"Error downloading validated file: {str(e)}")
            raise Http404("File not found")

    @action(detail=True, methods=['post'])
    def extract_orders(self, request, pk=None):
        """Extract orders from validated file synchronously"""
        try:
            validated_file = self.get_object()

            if validated_file.is_processed:
                return Response({
                    'error': 'Orders have already been extracted from this file'
                }, status=status.HTTP_400_BAD_REQUEST)

            # Direct synchronous call - no .delay()
            from .utils import extract_orders_from_validated_file_sync
            orders = extract_orders_from_validated_file_sync(validated_file)

            # Mark file as processed
            validated_file.is_processed = True
            validated_file.orders_extracted = len(orders)
            validated_file.save()

            logger.info(f"Successfully extracted {len(orders)} orders from {validated_file.file_name}")

            return Response({
                'success': True,
                'message': f'Successfully extracted {len(orders)} orders',
                'orders_extracted': len(orders),
                'orders': OrderListSerializer(orders, many=True).data
            })

        except Exception as e:
            logger.error(f"Error extracting orders from {validated_file.file_name}: {str(e)}")
            return Response({
                'error': f'Failed to extract orders: {str(e)}'
            }, status=status.HTTP_400_BAD_REQUEST)


class OrderViewSet(viewsets.ModelViewSet):
    """
    ViewSet for Order operations with file-specific filtering
    """
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['number', 'payment_reference_id', 'order_id']
    ordering_fields = ['created_at', 'amount', 'status', 'payment_status']
    ordering = ['-created_at']
    pagination_class = StandardResultsSetPagination

    def get_queryset(self):
        queryset = Order.objects.filter(validated_file__user=self.request.user)

        # Filter by validated file ID if provided
        validated_file_id = self.request.query_params.get('validated_file')
        if validated_file_id:
            queryset = queryset.filter(validated_file__id=validated_file_id)

        # Filter by status if provided
        status_filter = self.request.query_params.get('status')
        if status_filter:
            queryset = queryset.filter(status=status_filter)

        # Filter by payment status if provided
        payment_status = self.request.query_params.get('payment_status')
        if payment_status:
            queryset = queryset.filter(payment_status=payment_status)

        return queryset.select_related('validated_file')

    def get_serializer_class(self):
        if self.action == 'list':
            return OrderListSerializer
        return OrderSerializer

    def list(self, request, *args, **kwargs):
        """Enhanced list with file-specific stats"""
        queryset = self.filter_queryset(self.get_queryset())

        # Get file-specific statistics if filtering by file
        validated_file_id = request.query_params.get('validated_file')
        stats = {}

        if validated_file_id:
            try:
                validated_file = ValidatedFile.objects.get(
                    id=validated_file_id,
                    user=request.user
                )

                file_orders = queryset.filter(validated_file=validated_file)
                stats = {
                    'file_name': validated_file.file_name,
                    'total_orders': file_orders.count(),
                    'sent': file_orders.filter(status='sent').count(),
                    'delivered': file_orders.filter(status='delivered').count(),
                    'read': file_orders.filter(status='read').count(),
                    'failed': file_orders.filter(status='failed').count(),
                    'pending': file_orders.filter(status='pending').count(),
                    'payment_completed': file_orders.filter(payment_status='completed').count(),
                    'payment_pending': file_orders.filter(payment_status='pending').count(),
                    'payment_failed': file_orders.filter(payment_status='failed').count(),
                }
            except ValidatedFile.DoesNotExist:
                pass

        page = self.paginate_queryset(queryset)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            response = self.get_paginated_response(serializer.data)
            if stats:
                response.data['file_stats'] = stats
            return response

        serializer = self.get_serializer(queryset, many=True)
        return Response({
            'results': serializer.data,
            'file_stats': stats
        })

    @action(detail=False, methods=['get'])
    def refresh_status(self, request):
        """Refresh and return latest status for orders"""
        validated_file_id = request.query_params.get('validated_file')
        if not validated_file_id:
            return Response({
                'error': 'validated_file parameter required'
            }, status=status.HTTP_400_BAD_REQUEST)

        # Return fresh data
        return self.list(request)

    @action(detail=False, methods=['get'])
    def export_csv(self, request):
        """Export orders as CSV"""
        import csv
        from django.http import HttpResponse

        validated_file_id = request.query_params.get('validated_file')
        if not validated_file_id:
            return Response({
                'error': 'validated_file parameter required'
            }, status=status.HTTP_400_BAD_REQUEST)

        try:
            validated_file = ValidatedFile.objects.get(
                id=validated_file_id,
                user=request.user
            )

            orders = Order.objects.filter(
                validated_file=validated_file
            ).order_by('-created_at')

            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = f'attachment; filename="orders_{validated_file.file_name}_{timezone.now().strftime("%Y%m%d_%H%M%S")}.csv"'

            writer = csv.writer(response)
            writer.writerow([
                'Order ID', 'Phone Number', 'Order Items', 'Amount',
                'Message Status', 'Payment Status', 'Sent At', 'Delivered At',
                'Read At', 'Payment Method', 'Amount Captured', 'Created At'
            ])

            for order in orders:
                writer.writerow([
                    order.order_id,
                    order.number,
                    order.order_items,
                    order.amount,
                    order.status,
                    order.payment_status,
                    order.sent_at.strftime('%Y-%m-%d %H:%M:%S') if order.sent_at else '',
                    order.delivered_at.strftime('%Y-%m-%d %H:%M:%S') if order.delivered_at else '',
                    order.read_at.strftime('%Y-%m-%d %H:%M:%S') if order.read_at else '',
                    order.payment_method or '',
                    str(order.payment_amount_captured) if order.payment_amount_captured else '',
                    order.created_at.strftime('%Y-%m-%d %H:%M:%S')
                ])

            return response

        except ValidatedFile.DoesNotExist:
            return Response({
                'error': 'Validated file not found'
            }, status=status.HTTP_404_NOT_FOUND)

    @action(detail=False, methods=['get'])
    def export_excel(self, request):
        """Export orders as Excel"""
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from django.http import HttpResponse
        import io

        validated_file_id = request.query_params.get('validated_file')
        if not validated_file_id:
            return Response({
                'error': 'validated_file parameter required'
            }, status=status.HTTP_400_BAD_REQUEST)

        try:
            validated_file = ValidatedFile.objects.get(
                id=validated_file_id,
                user=request.user
            )

            orders = Order.objects.filter(
                validated_file=validated_file
            ).order_by('-created_at')

            # Create workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Orders"

            # Headers
            headers = [
                'Order ID', 'Phone Number', 'Order Items', 'Amount',
                'Message Status', 'Payment Status', 'Sent At', 'Delivered At',
                'Read At', 'Payment Method', 'Amount Captured', 'Created At'
            ]

            # Style headers
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill

            # Data rows
            for row, order in enumerate(orders, 2):
                ws.cell(row=row, column=1, value=order.order_id)
                ws.cell(row=row, column=2, value=order.number)
                ws.cell(row=row, column=3, value=order.order_items)
                ws.cell(row=row, column=4, value=order.amount)
                ws.cell(row=row, column=5, value=order.status)
                ws.cell(row=row, column=6, value=order.payment_status)
                ws.cell(row=row, column=7, value=order.sent_at.strftime('%Y-%m-%d %H:%M:%S') if order.sent_at else '')
                ws.cell(row=row, column=8, value=order.delivered_at.strftime('%Y-%m-%d %H:%M:%S') if order.delivered_at else '')
                ws.cell(row=row, column=9, value=order.read_at.strftime('%Y-%m-%d %H:%M:%S') if order.read_at else '')
                ws.cell(row=row, column=10, value=order.payment_method or '')
                ws.cell(row=row, column=11, value=str(order.payment_amount_captured) if order.payment_amount_captured else '')
                ws.cell(row=row, column=12, value=order.created_at.strftime('%Y-%m-%d %H:%M:%S'))

            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            # Save to bytes
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            response = HttpResponse(
                output.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="orders_{validated_file.file_name}_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'

            return response

        except ValidatedFile.DoesNotExist:
            return Response({
                'error': 'Validated file not found'
            }, status=status.HTTP_404_NOT_FOUND)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def processing_stats(request):
    """
    Get processing statistics for current user
    """
    user = request.user

    stats = {
        'total_files_uploaded': ChatFile.objects.filter(user=user).count(),
        'total_files_processed': ParsedChatFile.objects.filter(user=user).count(),
        'total_files_validated': ValidatedFile.objects.filter(user=user).count(),
        'total_orders_extracted': Order.objects.filter(validated_file__user=user).count(),
        'total_messages_sent': Order.objects.filter(validated_file__user=user, status__in=['sent', 'delivered', 'read']).count(),
    }

    # Calculate processing success rate
    if stats['total_files_uploaded'] > 0:
        stats['processing_success_rate'] = (stats['total_files_processed'] / stats['total_files_uploaded']) * 100
    else:
        stats['processing_success_rate'] = 0.0

    serializer = ProcessingStatsSerializer(stats)
    return Response(serializer.data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def analytics_dashboard(request):
    """
    Comprehensive analytics dashboard with file-wise breakdown and payment tracking
    """
    user = request.user

    # Overall statistics
    total_orders = Order.objects.filter(validated_file__user=user).count()
    total_files = ValidatedFile.objects.filter(user=user).count()

    # Message status breakdown
    message_stats = {
        'total_orders': total_orders,
        'sent': Order.objects.filter(validated_file__user=user, status='sent').count(),
        'delivered': Order.objects.filter(validated_file__user=user, status='delivered').count(),
        'read': Order.objects.filter(validated_file__user=user, status='read').count(),
        'failed': Order.objects.filter(validated_file__user=user, status='failed').count(),
        'pending': Order.objects.filter(validated_file__user=user, status='pending').count(),
    }

    # Payment statistics
    payment_stats = {
        'completed': Order.objects.filter(validated_file__user=user, payment_status='completed').count(),
        'pending': Order.objects.filter(validated_file__user=user, payment_status='pending').count(),
        'failed': Order.objects.filter(validated_file__user=user, payment_status='failed').count(),
        'initiated': Order.objects.filter(validated_file__user=user, payment_status='initiated').count(),
    }

    # Calculate conversion rates
    message_success_rate = ((message_stats['sent'] + message_stats['delivered'] + message_stats['read']) / total_orders * 100) if total_orders > 0 else 0
    payment_conversion_rate = (payment_stats['completed'] / total_orders * 100) if total_orders > 0 else 0
    read_rate = (message_stats['read'] / total_orders * 100) if total_orders > 0 else 0

    # Revenue calculation
    total_revenue = Order.objects.filter(
        validated_file__user=user,
        payment_status='completed'
    ).aggregate(
        total=Sum('payment_amount_captured')
    )['total'] or 0

    # File-wise breakdown
    file_breakdown = []
    for validated_file in ValidatedFile.objects.filter(user=user).order_by('-uploaded_at'):
        file_orders = Order.objects.filter(validated_file=validated_file)
        file_total = file_orders.count()

        if file_total > 0:
            file_stats = {
                'file_id': str(validated_file.id),
                'file_name': validated_file.file_name,
                'uploaded_at': validated_file.uploaded_at,
                'total_orders': file_total,
                'sent': file_orders.filter(status='sent').count(),
                'delivered': file_orders.filter(status='delivered').count(),
                'read': file_orders.filter(status='read').count(),
                'failed': file_orders.filter(status='failed').count(),
                'pending': file_orders.filter(status='pending').count(),
                'payment_completed': file_orders.filter(payment_status='completed').count(),
                'payment_pending': file_orders.filter(payment_status='pending').count(),
                'payment_failed': file_orders.filter(payment_status='failed').count(),
                'success_rate': ((file_orders.filter(status__in=['sent', 'delivered', 'read']).count()) / file_total * 100),
                'payment_conversion': (file_orders.filter(payment_status='completed').count() / file_total * 100),
                'revenue': file_orders.filter(payment_status='completed').aggregate(total=Sum('payment_amount_captured'))['total'] or 0
            }
            file_breakdown.append(file_stats)

    # Recent activity (last 10 orders)
    recent_orders = Order.objects.filter(
        validated_file__user=user
    ).select_related('validated_file').order_by('-updated_at')[:10]

    recent_activity = []
    for order in recent_orders:
        recent_activity.append({
            'order_id': order.order_id,
            'phone': order.number,
            'status': order.status,
            'payment_status': order.payment_status,
            'file_name': order.validated_file.file_name,
            'updated_at': order.updated_at
        })

    return Response({
        'overview': {
            'total_orders': total_orders,
            'total_files': total_files,
            'message_success_rate': round(message_success_rate, 2),
            'payment_conversion_rate': round(payment_conversion_rate, 2),
            'read_rate': round(read_rate, 2),
            'total_revenue': float(total_revenue)
        },
        'message_stats': message_stats,
        'payment_stats': payment_stats,
        'file_breakdown': file_breakdown,
        'recent_activity': recent_activity
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def process_stream(request, file_id):
    """
    Stream processing progress for a chat file
    """
    def stream_progress():
        try:
            # Check if file exists and belongs to user
            chat_file = ChatFile.objects.get(
                id=file_id,
                user=request.user
            )

            # Send initial status
            yield f"data: {json.dumps({'status': 'starting', 'message': 'Initializing processing...', 'progress': 0})}\n\n"

            # Monitor processing progress
            max_wait_time = 600  # 10 minutes
            check_interval = 2   # Check every 2 seconds
            elapsed_time = 0

            while elapsed_time < max_wait_time:
                # Refresh from database
                chat_file.refresh_from_db()

                if chat_file.processing_error:
                    yield f"data: {json.dumps({'status': 'error', 'message': chat_file.processing_error, 'progress': 0})}\n\n"
                    break

                if chat_file.is_processed:
                    # Check if processed file was created
                    try:
                        processed_file = chat_file.processed_file
                        yield f"data: {json.dumps({'status': 'completed', 'message': 'Processing completed successfully!', 'progress': 100, 'processed_file_id': str(processed_file.id)})}\n\n"
                    except ParsedChatFile.DoesNotExist:
                        yield f"data: {json.dumps({'status': 'error', 'message': 'Processing completed but no output file found', 'progress': 0})}\n\n"
                    break

                # Calculate rough progress based on elapsed time
                # This is a simple estimation - you could make this more sophisticated
                estimated_progress = min(int((elapsed_time / 300) * 80), 80)  # Cap at 80% until completion

                if elapsed_time < 30:
                    message = "Parsing chat messages..."
                    progress = estimated_progress
                elif elapsed_time < 120:
                    message = "Classifying messages with AI..."
                    progress = min(estimated_progress + 10, 85)
                else:
                    message = "Generating final output file..."
                    progress = min(estimated_progress + 15, 95)

                yield f"data: {json.dumps({'status': 'processing', 'message': message, 'progress': progress})}\n\n"

                time.sleep(check_interval)
                elapsed_time += check_interval

            # If we exit the loop without completion, it's a timeout
            if not chat_file.is_processed and not chat_file.processing_error:
                yield f"data: {json.dumps({'status': 'timeout', 'message': 'Processing is taking longer than expected. Please check back later.', 'progress': 0})}\n\n"

        except ChatFile.DoesNotExist:
            yield f"data: {json.dumps({'status': 'error', 'message': 'File not found or access denied', 'progress': 0})}\n\n"
        except Exception as e:
            logger.error(f"Error in process stream: {str(e)}")
            yield f"data: {json.dumps({'status': 'error', 'message': 'An unexpected error occurred', 'progress': 0})}\n\n"

    response = StreamingHttpResponse(
        stream_progress(),
        content_type='text/event-stream'
    )
    response['Cache-Control'] = 'no-cache'
    response['Connection'] = 'keep-alive'
    response['Access-Control-Allow-Origin'] = '*'
    response['Access-Control-Allow-Headers'] = 'Cache-Control'

    return response


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def send_order_messages(request):
    """
    Send WhatsApp messages for orders synchronously with 2-second delay

    Expected payload:
    {
        "validated_file_id": "uuid-here"
    }

    Returns:
    {
        "success": true,
        "total_orders": 5,
        "sent_count": 4,
        "failed_count": 1,
        "total_time_seconds": 10,
        "results": [
            {
                "order_id": "2025091500001",
                "number": "919474816594",
                "status": "sent",
                "message_id": "wamid_xyz"
            }
        ]
    }
    """
    try:
        validated_file_id = request.data.get('validated_file_id')

        if not validated_file_id:
            return Response({
                'error': 'validated_file_id is required'
            }, status=status.HTTP_400_BAD_REQUEST)

        # Get validated file and verify ownership
        try:
            validated_file = ValidatedFile.objects.get(
                id=validated_file_id,
                user=request.user
            )
        except ValidatedFile.DoesNotExist:
            return Response({
                'error': 'Validated file not found or access denied'
            }, status=status.HTTP_404_NOT_FOUND)

        # Get pending orders for this validated file
        orders = Order.objects.filter(
            validated_file=validated_file,
            status='pending'
        ).order_by('created_at')

        if not orders.exists():
            return Response({
                'error': 'No pending orders found for this validated file'
            }, status=status.HTTP_400_BAD_REQUEST)

        logger.info(f"Starting to send {orders.count()} order messages for {validated_file.file_name}")

        # Initialize WhatsApp service
        whatsapp_service = OrderWhatsAppService()

        # Send messages with 2-second delay
        results = whatsapp_service.send_bulk_orders(orders, delay_seconds=2.0)

        # Return detailed results
        return Response({
            'success': True,
            'message': f'Processed {orders.count()} orders',
            'total_orders': orders.count(),
            'sent_count': results['sent_count'],
            'failed_count': results['failed_count'],
            'total_time_seconds': results['total_time_seconds'],
            'results': results['details']
        })

    except Exception as e:
        logger.error(f"Error in send_order_messages: {str(e)}")
        return Response({
            'error': f'Failed to send messages: {str(e)}'
        }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
